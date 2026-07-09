"""Low-level drawing helpers built on top of python-pptx."""
from __future__ import annotations

from typing import Iterable, Optional, Sequence

from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from . import theme


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


def no_line(shape) -> None:
    shape.line.fill.background()


def solid_fill(shape, hexstr: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(hexstr)


def add_background(slide, hexstr: str = theme.WHITE):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, theme.SLIDE_W, theme.SLIDE_H)
    solid_fill(rect, hexstr)
    no_line(rect)
    rect.shadow.inherit = False
    # send to back
    spTree = slide.shapes._spTree
    spTree.remove(rect._element)
    spTree.insert(2, rect._element)
    return rect


def add_rect(slide, left, top, width, height, fill=None, line=None, line_w=Pt(0.75), shadow=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if fill is not None:
        solid_fill(shape, fill)
    else:
        shape.fill.background()
    if line is not None:
        shape.line.color.rgb = rgb(line)
        shape.line.width = line_w
    else:
        no_line(shape)
    return shape


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    size: int = 14,
    color: str = theme.INK,
    bold: bool = False,
    italic: bool = False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    font: str = theme.FONT,
    wrap: bool = True,
    line_spacing: Optional[float] = None,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = rgb(color)
    return box


def add_multirun_text(slide, left, top, width, height, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    """runs: list of (text, size, color, bold) rendered on a single paragraph."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    for text, size, color, bold in runs:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = theme.FONT
        run.font.color.rgb = rgb(color)
    return box


def status_dot(slide, left, top, diameter, status_key: str):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    solid_fill(dot, theme.status_color(status_key))
    no_line(dot)
    dot.shadow.inherit = False
    return dot


def status_pill(slide, left, top, width, height, status_key: str):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    try:
        pill.adjustments[0] = 0.5
    except Exception:
        pass
    solid_fill(pill, theme.status_color(status_key))
    no_line(pill)
    pill.shadow.inherit = False
    tf = pill.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = theme.status_label(status_key)
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.name = theme.FONT
    run.font.color.rgb = rgb(theme.WHITE)
    return pill


def set_cell(cell, text, size=10, color=theme.INK, bold=False, align=PP_ALIGN.LEFT,
             fill=None, anchor=MSO_ANCHOR.MIDDLE):
    cell.margin_left = Pt(6)
    cell.margin_right = Pt(6)
    cell.margin_top = Pt(3)
    cell.margin_bottom = Pt(3)
    cell.vertical_anchor = anchor
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(theme.WHITE)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = ""
    run = p.add_run()
    run.text = "" if text is None else str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = theme.FONT
    run.font.color.rgb = rgb(color)


def add_table(slide, left, top, width, height, col_widths: Sequence[int], n_rows: int):
    n_cols = len(col_widths)
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = graphic_frame.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = w
    # kill the built-in style banding so our own fills show cleanly
    tbl_el = table._tbl
    tblPr = tbl_el.find(qn("a:tblPr"))
    if tblPr is not None:
        tblPr.set("firstRow", "0")
        tblPr.set("bandRow", "0")
    return table


def add_kpi_tile(slide, left, top, width, height, label, value, sublabel="", value_color=theme.PRIMARY):
    add_rect(slide, left, top, width, height, fill=theme.SURFACE_ALT)
    pad = Emu(137160)
    add_text(slide, left + pad, top + pad, width - 2 * pad, Emu(228600), label.upper(),
              size=10, color=theme.MUTED, bold=True)
    add_text(slide, left + pad, top + Emu(320000), width - 2 * pad, Emu(600000), value,
              size=26, color=value_color, bold=True)
    if sublabel:
        add_text(slide, left + pad, top + height - Emu(320000), width - 2 * pad, Emu(280000),
                  sublabel, size=9, color=theme.MUTED)


def fmt_num(n) -> str:
    try:
        f = float(n)
    except (TypeError, ValueError):
        return str(n) if n is not None else ""
    if f == int(f):
        return f"{int(f):,}"
    return f"{f:,.1f}"
