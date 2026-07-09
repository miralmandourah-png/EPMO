"""Reads target/committed constants out of the user's uploaded template.

Per the user's instruction: any number that is a BRI/benefit *target*
(committed BRI, strategic ambition, committed BRI by line of business,
committed savings) is effectively fixed for the year and unlikely to change
month to month -- so rather than have the user retype it, we read it once
from their own slides and seed the dashboard with it. Numbers that describe
*current status* against that target (percent on-track, actual BRI realized,
IPI/TI, milestone counts) are NOT extracted here -- those are exactly what
changes monthly and stay driven by manual entry or the other importers.

This is the read-direction mirror of pptx_export.py's write-direction
mapping, reusing the same coordinate-anchoring approach (position + non-
confidential structural labels), applied to the same calibrated deck.
"""
from __future__ import annotations

import re
from typing import Any, Dict, List, Optional, Tuple

from pptx import Presentation

from .pptx_export import (
    _text_boxes, _left_top_in, _find_near, _find_in_band, _fmt_int, EXEC_POS,
    LOB_INIT_CONFIG, LOB_PROJ_CONFIG, INIT_COLS, NAME_COL_L, PROJ_NAME_L,
    CX_PROJ_SLIDES, CX_PROJ_ROW_TS, PROJ_IPI_L, PROJ_TI_L,
)

EXEC_SUMMARY_SLIDE = 4
CONTEXT_SLIDE = 5

AMBITION_ROWS = [
    ("gwp", "GWP (SAR B)"),
    ("profit", "Net profit (SAR B)"),
    ("roe", "Return on avg equity"),
    ("initshare", "Initiative share of GWP"),
]
BENEFIT_LOBS = ["Health", "General", "Life", "Motor"]
GWP_LOB_NAMES = ["Health", "General Corp", "Motor", "Life", "General Retail"]
YEAR_LABELS = ["2026", "2027", "2028", "2029", "2030"]


def _num(text: str) -> Optional[float]:
    m = re.search(r"[\d][\d,]*\.?\d*", text or "")
    if not m:
        return None
    try:
        return float(m.group(0).replace(",", ""))
    except ValueError:
        return None


def _row_boxes(slide, label_text: str, tol_t: float = 0.15, left_max: Optional[float] = None,
                region_right: Optional[float] = None, region_left: Optional[float] = None):
    """All text boxes on the same visual row as the box whose text starts
    with label_text (case-insensitive), excluding the label box itself.

    `region_left`/`region_right` cap the horizontal band candidates may fall
    in, so an unrelated box in a different column region (e.g. the panel on
    the other side of the slide) at a coincidentally similar row height is
    never picked up.
    """
    label_box = None
    for sh in _text_boxes(slide):
        t = sh.text_frame.text.strip()
        if t.lower().startswith(label_text.lower()) and (left_max is None or _left_top_in(sh)[0] <= left_max):
            label_box = sh
            break
    if label_box is None:
        return None, []
    label_l, row_t = _left_top_in(label_box)
    others = []
    for sh in _text_boxes(slide):
        if sh._element is label_box._element:
            continue
        L, T = _left_top_in(sh)
        if L is None or abs(T - row_t) > tol_t:
            continue
        if region_right is not None and L > region_right:
            continue
        if region_left is not None and L < region_left:
            continue
        others.append((L, sh.text_frame.text.strip()))
    others.sort(key=lambda x: x[0])
    return label_box, others


def _extract_ambition(slide) -> Dict[str, Any]:
    updates = {}
    for key, label in AMBITION_ROWS:
        _, candidates = _row_boxes(slide, label, tol_t=0.4, left_max=1.0, region_right=4.6)
        v2026 = v2030 = None
        for _, text in candidates:
            if text.endswith("2026"):
                v2026 = text.rsplit(" ", 1)[0].strip()
            elif text.endswith("2030"):
                v2030 = text.rsplit(" ", 1)[0].strip()
        if v2026 is not None:
            updates[f"exec.amb.{key}_2026"] = v2026
        if v2030 is not None:
            updates[f"exec.amb.{key}_2030"] = v2030
    return updates


def _extract_benefit_by_lob(slide) -> List[Dict[str, Any]]:
    rows = []
    for lob in BENEFIT_LOBS:
        _, candidates = _row_boxes(slide, lob, tol_t=0.05, left_max=10.0, region_left=9.0)
        sar = None
        for _, text in candidates:
            m = re.search(r"([\d,]+)\s*m\b", text)
            if m:
                sar = float(m.group(1).replace(",", ""))
                break
        rows.append({"lob": lob, "sar": sar})
    return rows


def _extract_savings_breakdown(slide) -> List[Dict[str, Any]]:
    items = []
    for sh in _text_boxes(slide):
        t = sh.text_frame.text.strip()
        L, T = _left_top_in(sh)
        if L is None or L < 9.0 or L > 10.2 or T < 5.7:
            continue
        m = re.match(r"^(.+?)\s+([\d.]+(?:-[\d.]+)?)$", t)
        if not m:
            continue
        name, amount_raw = m.group(1).strip(), m.group(2)
        amount = _num(amount_raw)
        detail_box = _find_near(slide, L + 1.27, T, tol=0.05)
        note = detail_box.text_frame.text.strip() if detail_box else ""
        items.append({"name": name, "amount": amount, "note": note})
    return items


def _extract_gwp_years(slide) -> List[Dict[str, Any]]:
    rows = []
    for year in YEAR_LABELS:
        _, candidates = _row_boxes(slide, year, tol_t=0.1, left_max=0.7, region_right=5.9)
        nums = [c for c in candidates if re.match(r"^[\d,.]+$", c[1])]
        bau = _num(nums[0][1]) if len(nums) > 0 else None
        init = _num(nums[1][1]) if len(nums) > 1 else None
        rows.append({"year": year, "bau": bau, "init": init})
    return rows


def _extract_gwp_by_lob(slide) -> List[Dict[str, Any]]:
    rows = []
    for lob in GWP_LOB_NAMES:
        _, candidates = _row_boxes(slide, lob, tol_t=0.1, left_max=7.5, region_left=6.0)
        # exclude the "total" box, which sits at the row's rightmost position
        if not candidates:
            rows.append({"lob": lob, "bau": None, "init": None})
            continue
        without_total = candidates[:-1] if len(candidates) > 1 else candidates
        nums = [c for c in without_total if re.match(r"^[\d,.]+$", c[1])]
        bau = _num(nums[0][1]) if len(nums) > 0 else None
        init = _num(nums[1][1]) if len(nums) > 1 else None
        rows.append({"lob": lob, "bau": bau, "init": init})
    return rows


def _extract_lob_initiatives(prs) -> Dict[str, List[Dict[str, Any]]]:
    """Initiative name/note/committed-BRI are constant identifiers -- IPI,
    TI, BRI actual, status, and realization date are monthly and left alone."""
    out: Dict[str, List[Dict[str, Any]]] = {}
    for lob_id, cfg in LOB_INIT_CONFIG.items():
        slide_idx = cfg["slides"][0]
        if slide_idx > len(prs.slides):
            continue
        slide = prs.slides[slide_idx - 1]
        entries = []
        for i in range(cfg["n_rows"]):
            row_t = cfg["row0_t"] + i * cfg["row_h"]
            box = _find_near(slide, NAME_COL_L, row_t, tol=0.05)
            name = note = None
            if box:
                paras = box.text_frame.paragraphs
                if len(paras) >= 1:
                    name = paras[0].text.strip()
                if len(paras) >= 2:
                    note = paras[1].text.strip()
            if not name:
                continue
            bri_box = _find_in_band(slide, row_t, *INIT_COLS["bri_committed"])
            bri = _num(bri_box.text_frame.text) if bri_box else None
            entries.append({"name": name, "note": note or "", "bri_committed": bri})
        out[lob_id] = entries
    return out


def _extract_lob_projects(prs) -> Dict[str, List[Dict[str, Any]]]:
    """Project name is a constant identifier -- IPI/TI/comment are monthly."""
    out: Dict[str, List[Dict[str, Any]]] = {}
    for lob_id, cfg in LOB_PROJ_CONFIG.items():
        slide_idx = cfg["slides"][0]
        if slide_idx > len(prs.slides):
            continue
        slide = prs.slides[slide_idx - 1]
        entries = []
        for row_t in cfg["row_ts"]:
            box = _find_near(slide, PROJ_NAME_L, row_t, tol=0.05)
            name = box.text_frame.text.strip() if box else ""
            if name:
                entries.append({"project": name})
        out[lob_id] = entries
    return out


def _extract_cx_projects(prs) -> List[Dict[str, Any]]:
    if not CX_PROJ_SLIDES or CX_PROJ_SLIDES[0] > len(prs.slides):
        return []
    slide = prs.slides[CX_PROJ_SLIDES[0] - 1]
    entries = []
    for row_t in CX_PROJ_ROW_TS:
        box = _find_near(slide, PROJ_NAME_L, row_t, tol=0.05)
        name = box.text_frame.text.strip() if box else ""
        if name:
            entries.append({"name": name})
    return entries


def extract_constants(template_bytes: bytes) -> Dict[str, Any]:
    """Returns {"updates": {...scalars...}, "tables": {tableId: [rows]}}."""
    prs = Presentation(__import__("io").BytesIO(template_bytes))
    n = len(prs.slides)
    updates: Dict[str, Any] = {}
    tables: Dict[str, List[Dict[str, Any]]] = {}

    if EXEC_SUMMARY_SLIDE <= n:
        s4 = prs.slides[EXEC_SUMMARY_SLIDE - 1]
        box = _find_near(s4, *EXEC_POS["committed_bri"], tol=0.05)
        if box:
            v = _num(box.text_frame.text)
            if v is not None:
                updates["exec.bri.total"] = v

        for sh in _text_boxes(s4):
            t = sh.text_frame.text.strip()
            m = re.match(r"^SAVINGS BRI\s*—\s*SAR\s*([\d,]+)m$", t, re.I)
            if m:
                updates["exec.bri.savings_total"] = float(m.group(1).replace(",", ""))
                break

        updates.update(_extract_ambition(s4))
        tables["exec.benefit"] = _extract_benefit_by_lob(s4)
        tables["exec.savings"] = _extract_savings_breakdown(s4)

    if CONTEXT_SLIDE <= n:
        s5 = prs.slides[CONTEXT_SLIDE - 1]
        tables["ctx.gwp_year"] = _extract_gwp_years(s5)
        tables["ctx.gwp_lob"] = _extract_gwp_by_lob(s5)

    for lob_id, entries in _extract_lob_initiatives(prs).items():
        tables[f"{lob_id}.init"] = entries
    for lob_id, entries in _extract_lob_projects(prs).items():
        tables[f"{lob_id}.proj"] = entries
    cx_entries = _extract_cx_projects(prs)
    if cx_entries:
        tables["cx.proj"] = cx_entries

    return {"updates": updates, "tables": tables}
